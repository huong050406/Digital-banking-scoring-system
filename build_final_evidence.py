import pandas as pd
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment
)
from openpyxl.utils import get_column_letter


# ============================================================
# 1. ĐƯỜNG DẪN
# ============================================================

BASE_DIR = Path(r"D:\Ngân hàng số\Data giữa kì")

TIER1_FILE = BASE_DIR / "step3_tier1_results.xlsx"
TIER2_FILE = BASE_DIR / "step4_tier2_results.xlsx"

OUTPUT_FILE = BASE_DIR / "Evidence_Final.xlsx"


# ============================================================
# 2. THỨ TỰ CHUẨN
# ============================================================

BANK_ORDER = [
    "BIDV",
    "VCB",
    "OCB",
    "MB",
    "VietinBank"
]

CRITERION_ORDER = [
    "A1", "A2", "A3", "A4",
    "B1", "B2", "B3", "B4",
    "C1", "C2", "C3", "C4",
    "D1", "D2", "D3", "D4"
]

STAGE_ORDER = [
    "Tầng 1",
    "Tầng 2"
]


# ============================================================
# 3. ĐỌC FILE
# ============================================================

df_tier1 = pd.read_excel(TIER1_FILE)
df_tier2 = pd.read_excel(TIER2_FILE)

print("======================================")
print("BƯỚC 5 - TỔNG HỢP EVIDENCE")
print("======================================")

print("Tổng dòng Tầng 1 gốc:", len(df_tier1))
print("Tổng dòng Bước 4:", len(df_tier2))


# ============================================================
# 4. CHUẨN HÓA TÊN NGÂN HÀNG
# ============================================================

def normalize_bank(bank):

    if pd.isna(bank):
        return ""

    value = str(bank).strip()

    key = (
        value
        .lower()
        .replace(" ", "")
    )

    mapping = {
        "bidv": "BIDV",

        "vcb": "VCB",
        "vietcombank": "VCB",

        "ocb": "OCB",

        "mb": "MB",
        "mbbank": "MB",

        "vietinbank": "VietinBank",
        "ctg": "VietinBank"
    }

    return mapping.get(
        key,
        value
    )


df_tier1["bank"] = (
    df_tier1["bank"]
    .apply(normalize_bank)
)

df_tier2["bank"] = (
    df_tier2["bank"]
    .apply(normalize_bank)
)


# ============================================================
# 5. CHUẨN HÓA TIER2_PASS
# ============================================================

def normalize_bool(value):

    if isinstance(value, bool):
        return value

    if pd.isna(value):
        return False

    value = str(value).strip().lower()

    return value in {
        "true",
        "1",
        "yes",
        "y"
    }


if "tier2_pass" in df_tier2.columns:

    df_tier2["tier2_pass"] = (
        df_tier2["tier2_pass"]
        .apply(normalize_bool)
    )


df_tier2_pass = df_tier2[
    df_tier2["tier2_pass"] == True
].copy()


print(
    "Số dòng qua Tầng 2 trước khi gộp:",
    len(df_tier2_pass)
)


# ============================================================
# 6. CLEAN TEXT
# ============================================================

def clean_text(value):

    if pd.isna(value):
        return ""

    value = str(value)

    value = re.sub(
        r"\s+",
        " ",
        value
    )

    return value.strip()


for dataframe in [
    df_tier1,
    df_tier2,
    df_tier2_pass
]:

    if "context" in dataframe.columns:

        dataframe["context"] = (
            dataframe["context"]
            .apply(clean_text)
        )


# ============================================================
# 7. GỘP GIÁ TRỊ KHÔNG TRÙNG
# ============================================================

def merge_unique(series):

    values = []

    for value in series:

        if pd.isna(value):
            continue

        value = str(value).strip()

        if not value:
            continue

        if value not in values:
            values.append(value)

    return "; ".join(values)


# ============================================================
# 8. CHỌN STRENGTH CAO NHẤT
# ============================================================

STRENGTH_RANK = {
    "None": 0,
    "Weak": 1,
    "Medium": 2,
    "Strong": 3
}


def strongest_strength(series):

    best = "None"
    best_rank = 0

    for value in series:

        if pd.isna(value):
            continue

        value = str(value).strip()

        rank = STRENGTH_RANK.get(
            value,
            0
        )

        if rank > best_rank:
            best = value
            best_rank = rank

    return best


# ============================================================
# 9. KEY CỦA 1 EVIDENCE
# ============================================================

EVIDENCE_KEY = [
    "bank",
    "criterion",
    "criterion_name",
    "report_type",
    "filename",
    "page",
    "context"
]


# ============================================================
# 10. DEDUP TẦNG 1
# ============================================================

def deduplicate_tier1(df):

    working = df.copy()

    for col in [
        "match_type",
        "matched_keyword"
    ]:

        if col not in working.columns:
            working[col] = ""

    result = (
        working
        .groupby(
            EVIDENCE_KEY,
            dropna=False,
            as_index=False
        )
        .agg({
            "match_type":
                merge_unique,

            "matched_keyword":
                merge_unique
        })
    )

    return result


df_tier1_clean = deduplicate_tier1(
    df_tier1
)


print(
    "Tầng 1 sau deduplicate:",
    len(df_tier1_clean)
)

print(
    "Số dòng Tầng 1 đã gộp:",
    len(df_tier1) - len(df_tier1_clean)
)


# ============================================================
# 11. DEDUP TẦNG 2
# ============================================================

def deduplicate_tier2(df):

    working = df.copy()

    required_cols = [
        "match_type",
        "matched_keyword",
        "tier2_score",
        "tier2_strength",
        "tier2_signals",
        "tier2_pass",
        "tier2_rule"
    ]

    for col in required_cols:

        if col not in working.columns:
            working[col] = ""

    working["tier2_score"] = pd.to_numeric(
        working["tier2_score"],
        errors="coerce"
    ).fillna(0)

    result = (
        working
        .groupby(
            EVIDENCE_KEY,
            dropna=False,
            as_index=False
        )
        .agg({
            "match_type":
                merge_unique,

            "matched_keyword":
                merge_unique,

            "tier2_score":
                "max",

            "tier2_strength":
                strongest_strength,

            "tier2_signals":
                merge_unique,

            "tier2_pass":
                "max",

            "tier2_rule":
                merge_unique
        })
    )

    return result


df_tier2_clean = deduplicate_tier2(
    df_tier2_pass
)


print(
    "Tầng 2 sau deduplicate:",
    len(df_tier2_clean)
)

print(
    "Số dòng Tầng 2 đã gộp:",
    len(df_tier2_pass) - len(df_tier2_clean)
)


# ============================================================
# 12. THÊM FILTER_STAGE
# ============================================================

tier1_raw = df_tier1_clean.copy()

tier1_raw.insert(
    1,
    "filter_stage",
    "Tầng 1"
)


tier2_raw = df_tier2_clean.copy()

tier2_raw.insert(
    1,
    "filter_stage",
    "Tầng 2"
)


# ============================================================
# 13. ĐỒNG BỘ CỘT
# ============================================================

FINAL_COLUMNS = [
    "bank",
    "filter_stage",
    "criterion",
    "criterion_name",
    "report_type",
    "filename",
    "page",
    "match_type",
    "matched_keyword",
    "context",
    "tier2_score",
    "tier2_strength",
    "tier2_signals",
    "tier2_pass",
    "tier2_rule"
]


for col in FINAL_COLUMNS:

    if col not in tier1_raw.columns:
        tier1_raw[col] = ""

    if col not in tier2_raw.columns:
        tier2_raw[col] = ""


tier1_raw = tier1_raw[
    FINAL_COLUMNS
]

tier2_raw = tier2_raw[
    FINAL_COLUMNS
]


# ============================================================
# 14. TẠO EVIDENCE_RAW
# ============================================================

df_raw = pd.concat(
    [
        tier1_raw,
        tier2_raw
    ],
    ignore_index=True
)


# ============================================================
# 15. SORT EVIDENCE_RAW
# ============================================================

bank_rank = {
    bank: i
    for i, bank
    in enumerate(BANK_ORDER)
}

criterion_rank = {
    criterion: i
    for i, criterion
    in enumerate(CRITERION_ORDER)
}

stage_rank = {
    stage: i
    for i, stage
    in enumerate(STAGE_ORDER)
}


df_raw["_bank_order"] = (
    df_raw["bank"]
    .map(bank_rank)
    .fillna(999)
)

df_raw["_criterion_order"] = (
    df_raw["criterion"]
    .map(criterion_rank)
    .fillna(999)
)

df_raw["_stage_order"] = (
    df_raw["filter_stage"]
    .map(stage_rank)
    .fillna(999)
)


df_raw = df_raw.sort_values(
    by=[
        "_bank_order",
        "_stage_order",
        "_criterion_order",
        "page"
    ]
)


df_raw = df_raw.drop(
    columns=[
        "_bank_order",
        "_criterion_order",
        "_stage_order"
    ]
)


# ============================================================
# 16. TẠO EVIDENCE_BY_BANK
# ============================================================

presentation_rows = []


def blank_row():

    return {
        "row_type": "",
        "bank": "",
        "stage": "",
        "criterion": "",
        "criterion_name": "",
        "report_type": "",
        "filename": "",
        "page": "",
        "match_type": "",
        "matched_keyword": "",
        "context": "",
        "tier2_score": "",
        "tier2_strength": "",
        "tier2_signals": "",
        "tier2_rule": ""
    }


def add_bank_header(bank):

    row = blank_row()

    row["row_type"] = "BANK"
    row["bank"] = bank

    presentation_rows.append(row)


def add_stage_header(
    bank,
    stage
):

    row = blank_row()

    row["row_type"] = "STAGE"
    row["bank"] = bank
    row["stage"] = stage

    presentation_rows.append(row)


def add_empty_row(
    bank,
    stage
):

    row = blank_row()

    row["row_type"] = "EMPTY"
    row["bank"] = bank
    row["stage"] = stage
    row["context"] = "Không có evidence."

    presentation_rows.append(row)


def add_evidence_rows(
    bank,
    stage,
    dataframe
):

    if dataframe.empty:

        add_empty_row(
            bank,
            stage
        )

        return


    temp = dataframe.copy()

    temp["_criterion_order"] = (
        temp["criterion"]
        .map(criterion_rank)
        .fillna(999)
    )


    if stage == "Tầng 2":

        temp = temp.sort_values(
            by=[
                "_criterion_order",
                "tier2_score",
                "page"
            ],
            ascending=[
                True,
                False,
                True
            ]
        )

    else:

        temp = temp.sort_values(
            by=[
                "_criterion_order",
                "page"
            ]
        )


    for _, row_data in temp.iterrows():

        row = blank_row()

        row["row_type"] = "EVIDENCE"
        row["bank"] = bank
        row["stage"] = stage

        row["criterion"] = (
            row_data.get(
                "criterion",
                ""
            )
        )

        row["criterion_name"] = (
            row_data.get(
                "criterion_name",
                ""
            )
        )

        row["report_type"] = (
            row_data.get(
                "report_type",
                ""
            )
        )

        row["filename"] = (
            row_data.get(
                "filename",
                ""
            )
        )

        row["page"] = (
            row_data.get(
                "page",
                ""
            )
        )

        row["match_type"] = (
            row_data.get(
                "match_type",
                ""
            )
        )

        row["matched_keyword"] = (
            row_data.get(
                "matched_keyword",
                ""
            )
        )

        row["context"] = (
            row_data.get(
                "context",
                ""
            )
        )

        row["tier2_score"] = (
            row_data.get(
                "tier2_score",
                ""
            )
        )

        row["tier2_strength"] = (
            row_data.get(
                "tier2_strength",
                ""
            )
        )

        row["tier2_signals"] = (
            row_data.get(
                "tier2_signals",
                ""
            )
        )

        row["tier2_rule"] = (
            row_data.get(
                "tier2_rule",
                ""
            )
        )

        presentation_rows.append(
            row
        )


# ============================================================
# 17. BUILD THEO NGÂN HÀNG
# ============================================================

for bank in BANK_ORDER:

    add_bank_header(bank)


    # --------------------------------------------------------
    # TẦNG 1
    # --------------------------------------------------------

    add_stage_header(
        bank,
        "Tầng 1"
    )

    bank_tier1 = tier1_raw[
        tier1_raw["bank"] == bank
    ].copy()

    add_evidence_rows(
        bank,
        "Tầng 1",
        bank_tier1
    )


    # --------------------------------------------------------
    # TẦNG 2
    # --------------------------------------------------------

    add_stage_header(
        bank,
        "Tầng 2"
    )

    bank_tier2 = tier2_raw[
        tier2_raw["bank"] == bank
    ].copy()

    add_evidence_rows(
        bank,
        "Tầng 2",
        bank_tier2
    )


df_by_bank = pd.DataFrame(
    presentation_rows
)


# ============================================================
# 18. GHI 4 SHEET
# ============================================================

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    # Raw B3 nguyên bản
    df_tier1.to_excel(
        writer,
        sheet_name="Step3_Tier1",
        index=False
    )

    # Raw B4 nguyên bản
    df_tier2.to_excel(
        writer,
        sheet_name="Step4_Tier2",
        index=False
    )

    # Dataset sạch
    df_raw.to_excel(
        writer,
        sheet_name="Evidence_Raw",
        index=False
    )

    # Bản dễ đọc
    df_by_bank.to_excel(
        writer,
        sheet_name="Evidence_By_Bank",
        index=False
    )


# ============================================================
# 19. LOAD LẠI ĐỂ FORMAT
# ============================================================

wb = load_workbook(
    OUTPUT_FILE
)


# ============================================================
# 20. STYLE
# ============================================================

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)

bank_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

bank_font = Font(
    color="FFFFFF",
    bold=True,
    size=14
)

tier1_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

tier2_fill = PatternFill(
    fill_type="solid",
    fgColor="E2F0D9"
)

tier2_evidence_fill = PatternFill(
    fill_type="solid",
    fgColor="F2F8EE"
)

stage_font = Font(
    bold=True,
    size=12
)


# ============================================================
# 21. FORMAT 3 SHEET DỮ LIỆU
# ============================================================

for sheet_name in [
    "Step3_Tier1",
    "Step4_Tier2",
    "Evidence_Raw"
]:

    ws = wb[sheet_name]

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = (
        ws.dimensions
    )


    # Header
    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )


    # Width theo header
    for col_num in range(
        1,
        ws.max_column + 1
    ):

        header = ws.cell(
            row=1,
            column=col_num
        ).value

        letter = get_column_letter(
            col_num
        )

        if header == "context":

            ws.column_dimensions[
                letter
            ].width = 80

        elif header == "tier2_signals":

            ws.column_dimensions[
                letter
            ].width = 60

        elif header == "tier2_rule":

            ws.column_dimensions[
                letter
            ].width = 60

        elif header == "matched_keyword":

            ws.column_dimensions[
                letter
            ].width = 42

        elif header == "filename":

            ws.column_dimensions[
                letter
            ].width = 38

        elif header == "criterion_name":

            ws.column_dimensions[
                letter
            ].width = 28

        else:

            ws.column_dimensions[
                letter
            ].width = 18


# ============================================================
# 22. FORMAT EVIDENCE_BY_BANK
# ============================================================

ws = wb[
    "Evidence_By_Bank"
]

ws.freeze_panes = "A2"

max_row = ws.max_row
max_col = ws.max_column


# Header
for cell in ws[1]:

    cell.fill = header_fill
    cell.font = header_font

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )


headers = {
    cell.value: cell.column
    for cell in ws[1]
}


# ============================================================
# 23. FORMAT TỪNG ROW TYPE
# ============================================================

for row_num in range(
    2,
    max_row + 1
):

    row_type = ws.cell(
        row=row_num,
        column=headers["row_type"]
    ).value


    # --------------------------------------------------------
    # BANK HEADER
    # --------------------------------------------------------

    if row_type == "BANK":

        for col_num in range(
            1,
            max_col + 1
        ):

            cell = ws.cell(
                row=row_num,
                column=col_num
            )

            cell.fill = bank_fill
            cell.font = bank_font

        ws.row_dimensions[
            row_num
        ].height = 26


    # --------------------------------------------------------
    # STAGE HEADER
    # --------------------------------------------------------

    elif row_type == "STAGE":

        stage = ws.cell(
            row=row_num,
            column=headers["stage"]
        ).value

        fill = (
            tier1_fill
            if stage == "Tầng 1"
            else tier2_fill
        )

        for col_num in range(
            1,
            max_col + 1
        ):

            cell = ws.cell(
                row=row_num,
                column=col_num
            )

            cell.fill = fill
            cell.font = stage_font


    # --------------------------------------------------------
    # EVIDENCE
    # --------------------------------------------------------

    elif row_type == "EVIDENCE":

        stage = ws.cell(
            row=row_num,
            column=headers["stage"]
        ).value

        # Chỉ wrap cột dài
        for column_name in [
            "criterion_name",
            "filename",
            "matched_keyword",
            "context",
            "tier2_signals",
            "tier2_rule"
        ]:

            col_num = headers.get(
                column_name
            )

            if col_num:

                ws.cell(
                    row=row_num,
                    column=col_num
                ).alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )


        # Tầng 2 tô nhẹ
        if stage == "Tầng 2":

            for col_num in range(
                1,
                max_col + 1
            ):

                ws.cell(
                    row=row_num,
                    column=col_num
                ).fill = tier2_evidence_fill


    # --------------------------------------------------------
    # EMPTY
    # --------------------------------------------------------

    elif row_type == "EMPTY":

        for col_num in range(
            1,
            max_col + 1
        ):

            cell = ws.cell(
                row=row_num,
                column=col_num
            )

            cell.font = Font(
                italic=True,
                color="808080"
            )


# ============================================================
# 24. WIDTH EVIDENCE_BY_BANK
# ============================================================

WIDTH_MAP = {

    "row_type": 12,
    "bank": 14,
    "stage": 12,
    "criterion": 10,
    "criterion_name": 28,
    "report_type": 20,
    "filename": 38,
    "page": 9,
    "match_type": 22,
    "matched_keyword": 42,
    "context": 85,
    "tier2_score": 12,
    "tier2_strength": 16,
    "tier2_signals": 60,
    "tier2_rule": 60
}


for header_name, width in WIDTH_MAP.items():

    col_num = headers.get(
        header_name
    )

    if col_num:

        letter = get_column_letter(
            col_num
        )

        ws.column_dimensions[
            letter
        ].width = width


# ============================================================
# 25. FILTER
# ============================================================

ws.auto_filter.ref = (
    ws.dimensions
)


# ============================================================
# 26. LƯU
# ============================================================

wb.save(
    OUTPUT_FILE
)


# ============================================================
# 27. THỐNG KÊ CUỐI
# ============================================================

print("\n======================================")
print("KẾT QUẢ B5 SAU CLEAN")
print("======================================")

print(
    "Step3_Tier1 RAW:",
    len(df_tier1)
)

print(
    "Step4_Tier2 RAW:",
    len(df_tier2)
)

print(
    "Tier2 PASS RAW:",
    len(df_tier2_pass)
)

print(
    "Evidence Tầng 1 sau gộp:",
    len(tier1_raw)
)

print(
    "Evidence Tầng 2 sau gộp:",
    len(tier2_raw)
)

print(
    "Tổng Evidence_Raw:",
    len(df_raw)
)


print(
    "\n----- Evidence Tầng 1 theo ngân hàng -----"
)

print(
    tier1_raw
    .groupby(
        "bank",
        observed=True
    )
    .size()
)


print(
    "\n----- Evidence Tầng 2 theo ngân hàng -----"
)

print(
    tier2_raw
    .groupby(
        "bank",
        observed=True
    )
    .size()
)


print(
    "\n----- Evidence Tầng 2 theo tiêu chí -----"
)

print(
    tier2_raw
    .groupby(
        "criterion",
        observed=True
    )
    .size()
    .reindex(
        CRITERION_ORDER,
        fill_value=0
    )
)


print("\n======================================")
print("✅ HOÀN TẤT BƯỚC 5")
print("======================================")

print(
    "File:",
    OUTPUT_FILE
)

print("\n4 sheet:")
print("1. Step3_Tier1")
print("2. Step4_Tier2")
print("3. Evidence_Raw")
print("4. Evidence_By_Bank")